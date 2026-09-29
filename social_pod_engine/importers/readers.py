"""CSV/XLSX readers that normalize headers but never write to a database."""

from __future__ import annotations

import csv
import re
from collections.abc import Iterable, Mapping
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

_HEADER_SEPARATOR = re.compile(r"[\s-]+")
REQUIRED_COLUMNS = frozenset({"platform", "username"})


class ImportFileStructureError(ValueError):
    """A file-wide issue that must abort before individual row processing."""


def normalize_column_name(value: object) -> str:
    if not isinstance(value, str):
        return ""
    return _HEADER_SEPARATOR.sub("_", value.strip().lower())


def read_import_file(path: Path | str, *, sheet_name: str | None = None) -> list[dict[str, Any]]:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(file_path)
    if suffix == ".xlsx":
        return _read_xlsx(file_path, sheet_name=sheet_name)
    raise ImportFileStructureError("only .csv and .xlsx import files are supported")


def normalize_rows(rows: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    for row in rows:
        mapped: dict[str, Any] = {}
        for key, value in row.items():
            normalized_key = normalize_column_name(key)
            if not normalized_key:
                continue
            if normalized_key in mapped:
                raise ImportFileStructureError(f"duplicate normalized column: {normalized_key}")
            mapped[normalized_key] = value
        normalized.append(mapped)
    _validate_columns(normalized[0].keys() if normalized else ())
    return normalized


def _read_csv(path: Path) -> list[dict[str, Any]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            reader = csv.DictReader(source)
            if reader.fieldnames is None:
                raise ImportFileStructureError("CSV has no header row")
            _validate_columns(normalize_column_name(column) for column in reader.fieldnames)
            raw_rows = list(reader)
    except UnicodeDecodeError as exc:
        raise ImportFileStructureError("CSV must use UTF-8 encoding") from exc
    except OSError as exc:
        raise ImportFileStructureError(f"could not read CSV: {exc}") from exc
    return normalize_rows(raw_rows)


def _read_xlsx(path: Path, *, sheet_name: str | None) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - translate library-specific parse errors
        raise ImportFileStructureError(f"could not read XLSX: {exc}") from exc
    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise ImportFileStructureError(f"worksheet not found: {sheet_name}")
            sheet = workbook[sheet_name]
        else:
            sheet = next((item for item in workbook.worksheets if item.sheet_state == "visible"), None)
            if sheet is None:
                raise ImportFileStructureError("XLSX has no visible worksheet")
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ImportFileStructureError("XLSX is empty")
    header = rows[0]
    _validate_columns(normalize_column_name(column) for column in header)
    raw_rows = [dict(zip(header, values, strict=False)) for values in rows[1:] if any(value is not None for value in values)]
    return normalize_rows(raw_rows)


def _validate_columns(columns: Iterable[str]) -> None:
    available = set(columns)
    missing = REQUIRED_COLUMNS - available
    if missing:
        raise ImportFileStructureError(f"missing required columns: {', '.join(sorted(missing))}")
